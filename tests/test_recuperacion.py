"""Pruebas sin depender del motor OCR pesado ni de videos institucionales."""
import sys
from pathlib import Path
import tempfile
import unittest

import cv2
import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import agregar_etapa_marcada, fraccion_scroll_para_ancla, parsear_etapas, rutas_de_exportacion
from motor import (
    _rango_muestreo, _normalizar, leer_zona, procesar_video,
    ZonaLectura, ProcesamientoCancelado,
)

class OcrFalso:
    def read(self, image):
        return [("12.34", 0.99)]

class Tests(unittest.TestCase):
    def test_etapas(self):
        self.assertEqual(parsear_etapas('0-2, 3-'), [(0,2),(3,None)])
        self.assertEqual(agregar_etapa_marcada('',1,3),'1-3')
        for t in ('0-2, 2-4', '5-1', '1', '-2', '5-2, 7-8'):
            with self.assertRaises(ValueError, msg=t):
                parsear_etapas(t)

    def test_muestreo_inclusivo(self):
        self.assertEqual(_rango_muestreo(31,10,1,0,3), [0,10,20,30])
        self.assertEqual(_rango_muestreo(31,10,2,0,2.5), [0,20,25])
        self.assertEqual(_rango_muestreo(0,10,1),[])

    def test_zona_ocr(self):
        z=ZonaLectura('Velocidad',3,2,50,30)
        im=np.zeros((100,100,3),np.uint8)
        self.assertEqual(z.recortar(im).shape, (30,50,3))
        self.assertEqual(leer_zona(OcrFalso(), z.recortar(im), z).valor,'12.34')
        self.assertEqual(_normalizar('90:61',2),'90.61')
        self.assertIsNone(_normalizar('70',2))

    def test_exportacion_evidencia_y_etapas(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            video=root/'muestra.avi'
            writer=cv2.VideoWriter(str(video),cv2.VideoWriter_fourcc(*'MJPG'),10,(160,90))
            if not writer.isOpened():
                self.skipTest('No hay codificador MJPG')
            for i in range(40):
                f=np.zeros((90,160,3),np.uint8)
                cv2.putText(f, '12.34', (10,45),cv2.FONT_HERSHEY_SIMPLEX,1,(255,255,255),2)
                writer.write(f)
            writer.release()
            salida=root/'lecturas.xlsx'
            muestras=procesar_video(video, [ZonaLectura('Velocidad',0,10,110,50,2)],1,
                OcrFalso(),salida,root/'evidencia', intervalos=[(0,1),(2,3)])
            self.assertEqual([f['frame'] for f in muestras], [0,10,20,30])
            self.assertTrue(salida.exists())
            libro=load_workbook(salida)
            self.assertEqual(set(libro.sheetnames), {'etapa_1','detalle_etapa_1',
                'etapa_2','detalle_etapa_2','Resumen de secciones'})
            self.assertEqual(libro['etapa_1']['C2'].value,12.34)
            self.assertEqual(libro['etapa_2'].max_row,3)
            self.assertEqual(libro.sheetnames[-1], 'Resumen de secciones')
            self.assertEqual(list(libro['Resumen de secciones'].values), [
                ('Sección', 'Hoja de lecturas', 'Inicio (s)', 'Final (s)'),
                ('Sección 1', 'etapa_1', 0, 1),
                ('Sección 2', 'etapa_2', 2, 3),
            ])
            for nombre, inicio, fin in [('etapa_1', 0, 1), ('detalle_etapa_1', 0, 1),
                                         ('etapa_2', 2, 3), ('detalle_etapa_2', 2, 3)]:
                hoja = libro[nombre]
                self.assertEqual(hoja.max_row, 3)
                self.assertEqual(hoja.auto_filter.ref, 'A1:D3' if nombre.startswith('etapa_') else 'A1:I3')
            self.assertEqual(len(list((root/'evidencia').glob('*.jpg'))),12)
            self.assertTrue(list((root/'evidencia').glob('*.png')))

    def test_cancelar(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            with self.assertRaises(ValueError):
                procesar_video('no_existe.mp4', [ZonaLectura('a',0,0,10,10)],1,
                    OcrFalso(),root/'x.xlsx',root/'evidencia')

# Comprobaciones adicionales independientes del reconocimiento OCR comercial.
class TestsAuxiliares(unittest.TestCase):
    def test_deteccion_cambios_con_referencia_recortada(self):
        from deteccion_eventos_experimental import proponer_periodos_en_video
        with tempfile.TemporaryDirectory() as tmp:
            video=Path(tmp)/'senal.avi'
            writer=cv2.VideoWriter(str(video),cv2.VideoWriter_fourcc(*'MJPG'),10,(100,80))
            if not writer.isOpened():
                self.skipTest('No hay codificador MJPG')
            for i in range(40):
                f=np.zeros((80,100,3),np.uint8)
                if 10<=i<30:
                    f[20:60,20:60]=255
                writer.write(f)
            writer.release()
            ref=np.zeros((40,40,3),np.uint8)
            rangos=proponer_periodos_en_video(video,(20,20,40,40),ref,
                                                muestreo_s=.1,minimo_s=.5)
            self.assertTrue(rangos)
            self.assertAlmostEqual(rangos[0][0],1,delta=.2)
            self.assertAlmostEqual(rangos[0][1],2.9,delta=.2)

    def test_cancelacion_sin_excel_final(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            video=root/'cancelar.avi'
            writer=cv2.VideoWriter(str(video),cv2.VideoWriter_fourcc(*'MJPG'),10,(100,80))
            for _ in range(5):writer.write(np.zeros((80,100,3),np.uint8))
            writer.release()
            salida=root/'salida.xlsx'
            with self.assertRaises(ProcesamientoCancelado):
                procesar_video(video,[ZonaLectura('dato',0,0,20,20)], .1,
                    OcrFalso(),salida,root/'evidencia',cancelar=lambda:True)
            self.assertFalse(salida.exists())

if __name__ == '__main__':
    unittest.main()
